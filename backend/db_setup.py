"""
NutriSync Database Setup
=========================
Builds / upgrades nutrisync.db. Safe to run repeatedly: **user data is never
dropped** (the previous version dropped daily_logs on every rebuild, which
would have wiped a user's whole food history).

Tables:
  food_items        - 1,014 Indian foods (Anuvaad/IFCT) + data-quality grade + diet tag
  user_profile      - the Dynamic User Profile (one row, single-user app)
  daily_logs        - every meal ever logged (short-term = "today", long-term = history)
  detected_patterns - long-term memory: trends noticed across days (one row per type)
  food_preferences  - long-term memory: how often / when the user eats each food
  user_facts        - long-term memory: things the user told us (diet, dislikes...)

Usage:
  python backend/db_setup.py            # create/upgrade, keep all user data
  python backend/db_setup.py --reset    # wipe user data too (destructive!)
"""

import os
import sqlite3
import sys

import openpyxl

try:
    from backend import food_quality
    from backend.database import get_db_connection
except ImportError:  # run as a script
    import food_quality
    from database import get_db_connection

_DEFAULT_DB = (
    os.path.join(os.path.dirname(__file__), "..", "nutrisync.db")
    if os.path.exists(os.path.join(os.path.dirname(__file__), "..", "nutrisync.db"))
    else os.path.join(os.path.dirname(__file__), "nutrisync.db")
)
DB_PATH = os.getenv("NUTRISYNC_DB_PATH", _DEFAULT_DB)
EXCEL_PATH = (
    os.path.join(os.path.dirname(__file__), "data", "anuvaad.xlsx")
    if os.path.exists(os.path.join(os.path.dirname(__file__), "data", "anuvaad.xlsx"))
    else os.path.join(os.path.dirname(__file__), "..", "data", "anuvaad.xlsx")
)

SCHEMA = """
CREATE TABLE IF NOT EXISTS food_items (
    food_code                TEXT PRIMARY KEY,
    food_name                TEXT NOT NULL,
    energy_kcal_100g         REAL,
    carb_g_100g              REAL,
    protein_g_100g           REAL,
    fat_g_100g               REAL,
    fibre_g_100g             REAL,
    servings_unit            TEXT,
    unit_serving_energy_kcal REAL,
    unit_serving_carb_g      REAL,
    unit_serving_protein_g   REAL,
    unit_serving_fat_g       REAL,
    unit_serving_fibre_g     REAL,
    calcium_mg_100g          REAL,
    magnesium_mg_100g        REAL,
    sodium_mg_100g           REAL,
    potassium_mg_100g        REAL,
    iron_mg_100g             REAL,
    copper_mg_100g           REAL,
    zinc_mg_100g             REAL,
    vita_ug_100g             REAL,
    vite_mg_100g             REAL,
    vitd2_ug_100g            REAL,
    vitd3_ug_100g            REAL,
    vitk1_ug_100g            REAL,
    vitk2_ug_100g            REAL,
    folate_ug_100g           REAL,
    vitb1_mg_100g            REAL,
    vitb2_mg_100g            REAL,
    vitb3_mg_100g            REAL,
    vitb5_mg_100g            REAL,
    vitb6_mg_100g            REAL,
    vitb7_ug_100g            REAL,
    vitb9_ug_100g            REAL,
    vitc_mg_100g             REAL,
    unit_serving_calcium_mg  REAL,
    unit_serving_magnesium_mg REAL,
    unit_serving_sodium_mg   REAL,
    unit_serving_potassium_mg REAL,
    unit_serving_iron_mg     REAL,
    unit_serving_copper_mg   REAL,
    unit_serving_zinc_mg     REAL,
    unit_serving_vita_ug     REAL,
    unit_serving_vite_mg     REAL,
    unit_serving_vitd2_ug    REAL,
    unit_serving_vitd3_ug    REAL,
    unit_serving_vitk1_ug    REAL,
    unit_serving_vitk2_ug    REAL,
    unit_serving_folate_ug   REAL,
    unit_serving_vitb1_mg    REAL,
    unit_serving_vitb2_mg    REAL,
    unit_serving_vitb3_mg    REAL,
    unit_serving_vitb5_mg    REAL,
    unit_serving_vitb6_mg    REAL,
    unit_serving_vitb7_ug    REAL,
    unit_serving_vitb9_ug    REAL,
    unit_serving_vitc_mg     REAL
);

CREATE TABLE IF NOT EXISTS user_profile (
    id                  INTEGER PRIMARY KEY CHECK (id = 1),
    name                TEXT,
    age                 INTEGER,
    sex                 TEXT,
    height_cm           REAL,
    current_weight_kg   REAL,
    target_weight_kg    REAL,
    goal                TEXT,
    activity_level      TEXT,
    allergies           TEXT,
    medical_conditions  TEXT,
    sleep_schedule      TEXT,
    bmr_kcal            REAL,
    tdee_kcal           REAL,
    target_calories     REAL,
    target_protein_g    REAL,
    target_carbs_g      REAL,
    target_fat_g        REAL,
    target_water_l      REAL,
    onboarded_at        TEXT,
    photo_data          TEXT
);

CREATE TABLE IF NOT EXISTS daily_logs (
    log_id        INTEGER PRIMARY KEY AUTOINCREMENT,
    log_date      TEXT NOT NULL,
    log_time      TEXT NOT NULL,
    meal_type     TEXT,
    food_code     TEXT,
    food_name     TEXT,
    quantity      REAL,
    calories      REAL,
    protein_g     REAL,
    carbs_g       REAL,
    fat_g         REAL
);

CREATE TABLE IF NOT EXISTS detected_patterns (
    pattern_id    INTEGER PRIMARY KEY AUTOINCREMENT,
    detected_on   TEXT NOT NULL,
    pattern_type  TEXT,
    description   TEXT,
    still_active  INTEGER DEFAULT 1
);

-- Long-term memory: one row per food the user has ever logged.
CREATE TABLE IF NOT EXISTS food_preferences (
    food_code       TEXT PRIMARY KEY,
    food_name       TEXT,
    times_logged    INTEGER DEFAULT 0,
    first_eaten     TEXT,
    last_eaten      TEXT,
    breakfast_count INTEGER DEFAULT 0,
    lunch_count     INTEGER DEFAULT 0,
    dinner_count    INTEGER DEFAULT 0,
    snack_count     INTEGER DEFAULT 0,
    total_amount    REAL DEFAULT 0,
    grams_count     INTEGER DEFAULT 0
);

-- Long-term memory: explicit facts the user told us (diet, dislikes, notes...).
CREATE TABLE IF NOT EXISTS user_facts (
    fact_key    TEXT PRIMARY KEY,
    fact_value  TEXT NOT NULL,
    updated_at  TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS water_logs (
    log_id      INTEGER PRIMARY KEY AUTOINCREMENT,
    log_date    TEXT NOT NULL,
    log_time    TEXT NOT NULL,
    amount_l    REAL NOT NULL
);

CREATE TABLE IF NOT EXISTS chat_messages (
    msg_id      INTEGER PRIMARY KEY AUTOINCREMENT,
    role        TEXT NOT NULL,
    content     TEXT NOT NULL,
    tool_events TEXT,
    created_at  TEXT NOT NULL
);


CREATE TABLE IF NOT EXISTS reminders (
    id              INTEGER PRIMARY KEY AUTOINCREMENT,
    kind            TEXT NOT NULL,
    remind_time     TEXT NOT NULL,
    repeat          TEXT NOT NULL DEFAULT 'daily',
    once_date       TEXT,
    food_name       TEXT,
    quantity        REAL,
    unit            TEXT DEFAULT 'serving',
    meal_type       TEXT,
    water_l         REAL,
    auto_log        INTEGER NOT NULL DEFAULT 1,
    enabled         INTEGER NOT NULL DEFAULT 1,
    last_fired_date TEXT,
    created_at      TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS reminder_events (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    reminder_id INTEGER,
    fired_at    TEXT NOT NULL,
    title       TEXT NOT NULL,
    message     TEXT NOT NULL,
    logged      INTEGER NOT NULL DEFAULT 0,
    channels    TEXT DEFAULT ''
);

CREATE TABLE IF NOT EXISTS push_subscriptions (
    id              INTEGER PRIMARY KEY AUTOINCREMENT,
    endpoint        TEXT NOT NULL UNIQUE,
    p256dh          TEXT NOT NULL,
    auth            TEXT NOT NULL,
    user_agent      TEXT,
    created_at      TEXT NOT NULL,
    last_success_at TEXT
);

CREATE TABLE IF NOT EXISTS grocery_items (
    id         INTEGER PRIMARY KEY AUTOINCREMENT,
    name       TEXT NOT NULL,
    name_key   TEXT NOT NULL,
    quantity   REAL NOT NULL,
    unit       TEXT NOT NULL,
    added_at   TEXT NOT NULL,
    updated_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS custom_foods (
    id             INTEGER PRIMARY KEY AUTOINCREMENT,
    name           TEXT NOT NULL,
    name_key       TEXT NOT NULL UNIQUE,
    servings       REAL NOT NULL DEFAULT 1,
    serving_label  TEXT NOT NULL DEFAULT 'serving',
    total_grams    REAL,
    calories       REAL NOT NULL,
    protein_g      REAL NOT NULL,
    carbs_g        REAL NOT NULL,
    fat_g          REAL NOT NULL,
    diet_tag       TEXT,
    ingredients_json TEXT,
    notes          TEXT,
    created_at     TEXT NOT NULL,
    updated_at     TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS weight_logs (
    log_date   TEXT PRIMARY KEY,
    weight_kg  REAL NOT NULL,
    note       TEXT,
    logged_at  TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS daily_fitness (
    log_date         TEXT PRIMARY KEY,
    steps            INTEGER DEFAULT 0,
    calories_burned  REAL DEFAULT 0,
    running_minutes  REAL DEFAULT 0,
    distance_km      REAL,
    active_minutes   REAL,
    source           TEXT DEFAULT 'google_fit',
    synced_at        TEXT NOT NULL
);
"""

MICRONUTRIENT_COLS_100G = [
    ("calcium_mg", "calcium_mg_100g"),
    ("magnesium_mg", "magnesium_mg_100g"),
    ("sodium_mg", "sodium_mg_100g"),
    ("potassium_mg", "potassium_mg_100g"),
    ("iron_mg", "iron_mg_100g"),
    ("copper_mg", "copper_mg_100g"),
    ("zinc_mg", "zinc_mg_100g"),
    ("vita_ug", "vita_ug_100g"),
    ("vite_mg", "vite_mg_100g"),
    ("vitd2_ug", "vitd2_ug_100g"),
    ("vitd3_ug", "vitd3_ug_100g"),
    ("vitk1_ug", "vitk1_ug_100g"),
    ("vitk2_ug", "vitk2_ug_100g"),
    ("folate_ug", "folate_ug_100g"),
    ("vitb1_mg", "vitb1_mg_100g"),
    ("vitb2_mg", "vitb2_mg_100g"),
    ("vitb3_mg", "vitb3_mg_100g"),
    ("vitb5_mg", "vitb5_mg_100g"),
    ("vitb6_mg", "vitb6_mg_100g"),
    ("vitb7_ug", "vitb7_ug_100g"),
    ("vitb9_ug", "vitb9_ug_100g"),
    ("vitc_mg", "vitc_mg_100g"),
]

MICRONUTRIENT_SERVING_COLS = [
    ("unit_serving_calcium_mg", "unit_serving_calcium_mg"),
    ("unit_serving_magnesium_mg", "unit_serving_magnesium_mg"),
    ("unit_serving_sodium_mg", "unit_serving_sodium_mg"),
    ("unit_serving_potassium_mg", "unit_serving_potassium_mg"),
    ("unit_serving_iron_mg", "unit_serving_iron_mg"),
    ("unit_serving_copper_mg", "unit_serving_copper_mg"),
    ("unit_serving_zinc_mg", "unit_serving_zinc_mg"),
    ("unit_serving_vita_ug", "unit_serving_vita_ug"),
    ("unit_serving_vite_mg", "unit_serving_vite_mg"),
    ("unit_serving_vitd2_ug", "unit_serving_vitd2_ug"),
    ("unit_serving_vitd3_ug", "unit_serving_vitd3_ug"),
    ("unit_serving_vitk1_ug", "unit_serving_vitk1_ug"),
    ("unit_serving_vitk2_ug", "unit_serving_vitk2_ug"),
    ("unit_serving_folate_ug", "unit_serving_folate_ug"),
    ("unit_serving_vitb1_mg", "unit_serving_vitb1_mg"),
    ("unit_serving_vitb2_mg", "unit_serving_vitb2_mg"),
    ("unit_serving_vitb3_mg", "unit_serving_vitb3_mg"),
    ("unit_serving_vitb5_mg", "unit_serving_vitb5_mg"),
    ("unit_serving_vitb6_mg", "unit_serving_vitb6_mg"),
    ("unit_serving_vitb7_ug", "unit_serving_vitb7_ug"),
    ("unit_serving_vitb9_ug", "unit_serving_vitb9_ug"),
    ("unit_serving_vitc_mg", "unit_serving_vitc_mg"),
]

# Columns added after the first release. (table, column, DDL type)
MIGRATIONS = [
    ("food_items", "serving_grams", "REAL"),
    ("food_items", "quality", "TEXT"),
    ("food_items", "quality_note", "TEXT"),
    ("food_items", "diet_tag", "TEXT"),
    ("daily_logs", "unit", "TEXT DEFAULT 'serving'"),
    ("daily_logs", "serving_label", "TEXT"),
    ("user_profile", "diet", "TEXT DEFAULT 'any'"),
    ("user_profile", "photo_data", "TEXT"),
    ("daily_fitness", "extras", "TEXT"),
] + [
    ("food_items", target_col, "REAL")
    for _, target_col in MICRONUTRIENT_COLS_100G + MICRONUTRIENT_SERVING_COLS
]


def _columns(conn, table):
    return {r[1] for r in conn.execute(f"PRAGMA table_info({table})")}


def _safe_float(val):
    if val is None or val == "":
        return None
    try:
        return float(val)
    except (ValueError, TypeError):
        return None


def refresh_food_quality(conn) -> int:
    """(Re)grade every food row. Cheap (1k rows), idempotent."""
    conn.row_factory = sqlite3.Row
    rows = conn.execute("SELECT * FROM food_items").fetchall()
    updates = []
    for r in rows:
        q = food_quality.assess(r)
        updates.append((
            q["quality"],
            q["quality_note"],
            q["serving_grams"],
            food_quality.diet_tag(r["food_name"]),
            r["food_code"],
        ))
    chunk_size = 100
    for i in range(0, len(updates), chunk_size):
        conn.executemany(
            "UPDATE food_items SET quality=?, quality_note=?, serving_grams=?, diet_tag=? WHERE food_code=?",
            updates[i : i + chunk_size],
        )
    return len(rows)


def rebuild_food_preferences(conn) -> None:
    """Recompute long-term food habits from the raw log (used on migration)."""
    conn.execute("DELETE FROM food_preferences")
    conn.execute(
        """INSERT INTO food_preferences
           (food_code, food_name, times_logged, first_eaten, last_eaten,
            breakfast_count, lunch_count, dinner_count, snack_count, total_amount, grams_count)
           SELECT food_code, MAX(food_name), COUNT(*), MIN(log_date), MAX(log_date),
                  SUM(meal_type='breakfast'), SUM(meal_type='lunch'),
                  SUM(meal_type='dinner'), SUM(meal_type='snack'),
                  SUM(quantity), SUM(COALESCE(unit,'serving')='grams')
           FROM daily_logs WHERE food_code IS NOT NULL GROUP BY food_code"""
    )


def ensure_schema(db_path=None) -> None:
    """Create missing tables, add missing columns, dedupe patterns. Never drops data."""
    path = db_path or DB_PATH
    conn = get_db_connection(path)
    conn.executescript(SCHEMA)
    existing_cols = {}
    for table, column, ddl in MIGRATIONS:
        if table not in existing_cols:
            existing_cols[table] = _columns(conn, table)
        if column not in existing_cols[table]:
            conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {ddl}")
            existing_cols[table].add(column)

    # Old versions inserted a fresh pattern row on every chat message. Keep the
    # newest per type, then enforce one-row-per-type going forward.
    conn.execute(
        "DELETE FROM detected_patterns WHERE pattern_id NOT IN "
        "(SELECT MAX(pattern_id) FROM detected_patterns GROUP BY pattern_type)"
    )
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_pattern_type ON detected_patterns(pattern_type)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_logs_date ON daily_logs(log_date)")

    needs_grading = conn.execute(
        "SELECT COUNT(*) FROM food_items WHERE quality IS NULL"
    ).fetchone()[0]
    if needs_grading:
        refresh_food_quality(conn)

    conn.row_factory = None
    has_prefs = conn.execute("SELECT COUNT(*) FROM food_preferences").fetchone()[0]
    has_logs = conn.execute("SELECT COUNT(*) FROM daily_logs").fetchone()[0]
    if has_logs and not has_prefs:
        rebuild_food_preferences(conn)
    conn.commit()
    conn.close()


def load_foods(conn) -> int:
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    col = {name: i for i, name in enumerate(header)}
    
    col_names = [
        "food_code", "food_name", "energy_kcal_100g", "carb_g_100g", "protein_g_100g",
        "fat_g_100g", "fibre_g_100g", "servings_unit", "unit_serving_energy_kcal",
        "unit_serving_carb_g", "unit_serving_protein_g", "unit_serving_fat_g",
        "unit_serving_fibre_g",
    ] + [dst for _, dst in MICRONUTRIENT_COLS_100G] + [dst for _, dst in MICRONUTRIENT_SERVING_COLS]

    placeholders = ", ".join(["?"] * len(col_names))
    sql = f"INSERT OR REPLACE INTO food_items ({', '.join(col_names)}) VALUES ({placeholders})"

    batch = []
    for row in rows:
        if not row[col["food_name"]]:
            continue
        vals = [
            row[col["food_code"]],
            row[col["food_name"]],
            _safe_float(row[col.get("energy_kcal")]),
            _safe_float(row[col.get("carb_g")]),
            _safe_float(row[col.get("protein_g")]),
            _safe_float(row[col.get("fat_g")]),
            _safe_float(row[col.get("fibre_g")]),
            row[col.get("servings_unit")],
            _safe_float(row[col.get("unit_serving_energy_kcal")]),
            _safe_float(row[col.get("unit_serving_carb_g")]),
            _safe_float(row[col.get("unit_serving_protein_g")]),
            _safe_float(row[col.get("unit_serving_fat_g")]),
            _safe_float(row[col.get("unit_serving_fibre_g")]),
        ]
        for src, _ in MICRONUTRIENT_COLS_100G:
            idx = col.get(src)
            vals.append(_safe_float(row[idx]) if idx is not None else None)
        for src, _ in MICRONUTRIENT_SERVING_COLS:
            idx = col.get(src)
            vals.append(_safe_float(row[idx]) if idx is not None else None)
        batch.append(tuple(vals))

    chunk_size = 100
    for i in range(0, len(batch), chunk_size):
        conn.executemany(sql, batch[i : i + chunk_size])
    return len(batch)


def seed_if_needed() -> None:
    """Fast startup path (used by the Docker entrypoint): create missing tables and load the
    food table only when it is empty. Never touches your logs / profile / chat history, so it
    is safe to run on every boot -- including against Turso."""
    ensure_schema(DB_PATH)
    conn = get_db_connection(DB_PATH)
    count = conn.execute("SELECT COUNT(*) FROM food_items").fetchone()[0]
    if count == 0:
        inserted = load_foods(conn)
        refresh_food_quality(conn)
        conn.commit()
        print(f"Seeded {inserted} foods.")
    else:
        print(f"Database ready ({count} foods already loaded, user data untouched).")
    conn.close()


def build_database(reset_user_data: bool = False) -> None:
    """Create/upgrade the DB and (re)load the food table. User data is kept
    unless reset_user_data=True."""
    if reset_user_data:
        conn = get_db_connection(DB_PATH)
        conn.executescript(
            "DROP TABLE IF EXISTS user_profile; DROP TABLE IF EXISTS daily_logs; "
            "DROP TABLE IF EXISTS detected_patterns; DROP TABLE IF EXISTS food_preferences; "
            "DROP TABLE IF EXISTS user_facts; DROP TABLE IF EXISTS chat_messages; "
            "DROP TABLE IF EXISTS water_logs; DROP TABLE IF EXISTS reminders; "
            "DROP TABLE IF EXISTS reminder_events; DROP TABLE IF EXISTS grocery_items; "
            "DROP TABLE IF EXISTS custom_foods; DROP TABLE IF EXISTS weight_logs;"
        )
        conn.commit()
        conn.close()

    ensure_schema(DB_PATH)
    conn = get_db_connection(DB_PATH)
    inserted = load_foods(conn)
    refresh_food_quality(conn)
    conn.commit()
    graded = dict(conn.execute("SELECT quality, COUNT(*) FROM food_items GROUP BY quality").fetchall())
    conn.close()
    print(f"Database ready ({'Turso' if os.getenv('TURSO_DATABASE_URL') else DB_PATH}) -- {inserted} foods loaded. Data quality: {graded}")


if __name__ == "__main__":
    if "--if-needed" in sys.argv:
        seed_if_needed()
        sys.exit(0)
    if "--reset" in sys.argv:
        confirm = input("This DELETES all logged meals, profile and chat history. Type 'yes' to continue: ")
        if confirm.strip().lower() != "yes":
            sys.exit("Aborted.")
        build_database(reset_user_data=True)
    else:
        build_database()